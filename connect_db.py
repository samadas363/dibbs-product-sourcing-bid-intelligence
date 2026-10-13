
import os
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import create_engine, text
from sqlalchemy.engine import URL
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------
# 1. CONFIGURATION
# --------------------------------------------------

PROJECT_DIR = Path(__file__).resolve().parent
load_dotenv(PROJECT_DIR / ".env")

DB_HOST = os.getenv("DB_HOST", "localhost")
DB_PORT = int(os.getenv("DB_PORT", "5432"))
DB_NAME = os.getenv("DB_NAME", "supplier_quote_intelligence")
DB_USER = os.getenv("DB_USER", "postgres")
DB_PASSWORD = os.getenv("DB_PASSWORD")

if not DB_PASSWORD:
    raise ValueError("DB_PASSWORD is missing from your .env file.")

DATABASE_URL = URL.create(
    "postgresql+psycopg2",
    username=DB_USER,
    password=DB_PASSWORD,
    host=DB_HOST,
    port=DB_PORT,
    database=DB_NAME,
)

engine = create_engine(DATABASE_URL)

# --------------------------------------------------
# 2. READ SUPPLIER COMPARISON DATA
# --------------------------------------------------

try:
    with engine.connect() as connection:
        database, username = connection.execute(
            text("SELECT current_database(), current_user")
        ).fetchone()

    print("PostgreSQL connected successfully!")
    print(f"Database: {database}")
    print(f"User: {username}")

    query = text("""
        SELECT *
        FROM public.supplier_quote_comparison
        ORDER BY cost_rank NULLS LAST, supplier_name;
    """)

    with engine.connect() as connection:
        df = pd.read_sql(query, connection)

    if df.empty:
        print("No supplier comparison records found.")
    else:
        print("\nSupplier comparison:")
        print(df.to_string(index=False))

        # --------------------------------------------------
        # 3. PREPARE DATA FOR EXCEL
        # --------------------------------------------------

        for column in df.columns:
            if isinstance(df[column].dtype, pd.DatetimeTZDtype):
                # Excel does not support timezone-aware dates.
                df[column] = df[column].dt.tz_localize(None)

        output_file = PROJECT_DIR / "supplier_comparison_report.xlsx"

        # Find the lowest-cost eligible supplier for the summary.
        eligible = df[
            (df["quote_status"] == "ELIGIBLE - QUANTITY MET")
            & df["estimated_required_quantity_cost"].notna()
        ]

        if not eligible.empty:
            best_supplier = eligible.loc[
                eligible["estimated_required_quantity_cost"].idxmin()
            ]
            best_name = best_supplier["supplier_name"]
            best_cost = float(
                best_supplier["estimated_required_quantity_cost"]
            )
        else:
            best_name = "No eligible supplier"
            best_cost = None

        # --------------------------------------------------
        # 4. EXPORT AND FORMAT EXCEL
        # --------------------------------------------------

        with pd.ExcelWriter(
            output_file,
            engine="openpyxl",
            datetime_format="yyyy-mm-dd hh:mm"
        ) as writer:

            df.to_excel(
                writer,
                index=False,
                sheet_name="Supplier Comparison",
                startrow=4
            )

            ws = writer.sheets["Supplier Comparison"]

            # Report title and summary
            ws["A1"] = "DIBBS SUPPLIER QUOTE INTELLIGENCE"
            ws["A1"].font = Font(
                name="Arial", size=16, bold=True, color="FFFFFF"
            )
            ws["A1"].fill = PatternFill(
                "solid", fgColor="17365D"
            )
            ws.merge_cells(
                start_row=1, start_column=1,
                end_row=1, end_column=len(df.columns)
            )
            ws.row_dimensions[1].height = 30

            ws["A2"] = "Opportunity"
            ws["B2"] = (
                str(df["solicitation_number"].iloc[0])
                if "solicitation_number" in df.columns else "N/A"
            )
            ws["A3"] = "Lowest-cost eligible supplier"
            ws["B3"] = best_name
            ws["C3"] = "Estimated cost"
            ws["D3"] = best_cost

            for cell in ["A2", "A3", "C3"]:
                ws[cell].font = Font(bold=True, color="17365D")

            if best_cost is not None:
                ws["D3"].number_format = '$#,##0.00'

            # Header row is row 5
            header_row = 5
            for cell in ws[header_row]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="305496")
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

            ws.row_dimensions[header_row].height = 32
            ws.freeze_panes = "A6"
            ws.auto_filter.ref = ws.dimensions

            # Map column names to Excel column positions
            col_positions = {
                cell.value: cell.column
                for cell in ws[header_row]
            }

            # Currency columns
            currency_columns = [
                "unit_price",
                "shipping_cost",
                "quoted_total_cost",
                "estimated_required_quantity_cost",
            ]

            for name in currency_columns:
                if name in col_positions:
                    col_num = col_positions[name]
                    for row_num in range(header_row + 1, ws.max_row + 1):
                        ws.cell(row_num, col_num).number_format = '$#,##0.00'

            # Highlight eligibility statuses
            if "quote_status" in col_positions:
                status_col = col_positions["quote_status"]

                for row_num in range(header_row + 1, ws.max_row + 1):
                    status = ws.cell(row_num, status_col).value

                    if status == "ELIGIBLE - QUANTITY MET":
                        fill = PatternFill("solid", fgColor="E2F0D9")
                    elif status == "NOT ELIGIBLE - QUANTITY SHORTFALL":
                        fill = PatternFill("solid", fgColor="FCE4D6")
                    else:
                        fill = PatternFill("solid", fgColor="FFF2CC")

                    ws.cell(row_num, status_col).fill = fill
                    ws.cell(row_num, status_col).alignment = Alignment(
                        vertical="center",
                        wrap_text=True
                    )

            # Make columns readable
            for col_num in range(1, ws.max_column + 1):
                letter = get_column_letter(col_num)
                values = [
                    len(str(ws.cell(row_num, col_num).value))
                    if ws.cell(row_num, col_num).value is not None
                    else 0
                    for row_num in range(1, ws.max_row + 1)
                ]
                ws.column_dimensions[letter].width = min(
                    max(max(values) + 2, 12), 38
                )

            # Improve data alignment
            thin_gray = Side(style="thin", color="D9E2F3")
            for row in ws.iter_rows(
                min_row=header_row,
                max_row=ws.max_row,
                min_col=1,
                max_col=ws.max_column
            ):
                for cell in row:
                    cell.border = Border(bottom=thin_gray)
                    if cell.row > header_row:
                        cell.alignment = Alignment(
                            vertical="top",
                            wrap_text=True
                        )

        print("\nProfessional Excel report created successfully!")
        print(f"Report location: {output_file}")
        print(f"Rows exported: {len(df)}")
        print(f"Lowest-cost eligible supplier: {best_name}")

        if best_cost is not None:
            print(f"Estimated cost: ${best_cost:,.2f}")

except Exception as error:
    print(f"\nError: {error}")

finally:
    engine.dispose()
